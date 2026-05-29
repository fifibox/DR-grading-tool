import tkinter as tk
from tkinter import filedialog, messagebox
import os
import csv
from pathlib import Path
from PIL import Image, ImageTk, ImageEnhance
from openpyxl import Workbook, load_workbook

# DR grade definitions
DR_GRADES = [
    {"label": "-1", "name": "Ungradable", "color": "#030304", "desc": "Unable to grade"},
    {"label": "0", "name": "No DR",        "color": "#1a7f37", "desc": "No apparent diabetic retinopathy"},
    {"label": "1", "name": "Mild",          "color": "#e09c14", "desc": "Microaneurysms only"},
    {"label": "2", "name": "Moderate",      "color": "#d95c09", "desc": "More than just microaneurysms but less than severe"},
    {"label": "3", "name": "Severe",        "color": "#cf222e", "desc": "20+ hemorrhages, venous beading, or IRMA"},
    {"label": "4", "name": "Proliferative", "color": "#8250df", "desc": "Neovascularisation or vitreous/pre-retinal hemorrhage"},
]

SUPPORTED_EXTS = {".jpg", ".jpeg", ".png", ".bmp", ".tiff", ".tif", ".webp"}

# Colours
BG        = "#ffffff"
BG_SIDE   = "#f6f8fa"
BG_SEL    = "#0969da"
FG        = "#1f2328"
FG_MUTED  = "#656d76"
BORDER    = "#d0d7de"


class DRGrader(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Diabetic Retinopathy Grader")
        self.configure(bg=BG)
        self.minsize(1100, 700)

        # State
        self.image_folder  = None
        self.image_files   = []
        self.current_index = -1
        self.grades        = {}
        self.current_photo = None
        self.selected_grade = tk.StringVar(value="")

        # CSV state
        self.csv_path      = None
        self.csv_data      = {}  # {filename: {"va": value, "grader_initial": "", "label": "", "note": ""}}
        self.use_csv_mode  = False
        self.grader_initial_var = tk.StringVar(value="")
        self.note_var = tk.StringVar(value="")

        # Image processing state
        self.brightness = 1.0
        self.red_free_filter = False

        # Zoom / pan state
        self._pil_image   = None
        self._zoom        = 1.0
        self._pan_x       = 0.0
        self._pan_y       = 0.0
        self._drag_start  = None
        self._drag_pan    = None

        self._build_ui()
        self._bind_keys()

    # ------------------------------------------------------------------ #
    #  UI                                                                  #
    # ------------------------------------------------------------------ #
    def _build_ui(self):
        # ── Top bar ──────────────────────────────────────────────────────
        topbar = tk.Frame(self, bg=BG_SIDE, pady=8,
                          highlightthickness=1, highlightbackground=BORDER)
        topbar.pack(fill="x", side="top")


        self.folder_label = tk.Label(topbar, text="No folder selected",
                                     font=("Helvetica", 15), fg=FG_MUTED, bg=BG_SIDE)
        self.folder_label.pack(side="left", padx=8)

        btn = {"font": ("Helvetica", 13), "relief": "solid", "bd": 1,
               "bg": BG, "fg": FG, "activebackground": BG, "activeforeground": FG,
               "cursor": "hand2", "pady": 4, "padx": 12, "highlightthickness": 0}

        tk.Button(topbar, text="Open Folder", command=self._open_folder,
                  **btn).pack(side="left", padx=6)
        
        tk.Button(topbar, text="Open Excel", command=self._open_csv,
                  **btn).pack(side="left", padx=16)

        tk.Label(topbar, text="Grader initial:", font=("Helvetica", 13),
                 fg=FG, bg=BG_SIDE).pack(side="left", padx=(0, 6))

        self.grader_entry = tk.Entry(topbar, textvariable=self.grader_initial_var,
                                     font=("Helvetica", 13), width=6)
        self.grader_entry.pack(side="left", padx=(0, 16))
        self.grader_entry.bind("<FocusOut>", lambda e: self._on_grader_changed())

        tk.Button(topbar, text="Save to Excel", command=self._save_csv,
                  **btn).pack(side="right", padx=16)

        self.progress_label = tk.Label(topbar, text="0 / 0  graded",
                                       font=("Helvetica", 15), fg=FG_MUTED, bg=BG_SIDE)
        self.progress_label.pack(side="right", padx=8)

        # ── Main content ─────────────────────────────────────────────────
        content = tk.Frame(self, bg=BG)
        content.pack(fill="both", expand=True)

        # Left panel – file list
        left = tk.Frame(content, bg=BG_SIDE, width=210,
                        highlightthickness=1, highlightbackground=BORDER)
        left.pack(side="left", fill="y")
        left.pack_propagate(False)

        tk.Label(left, text="Images", font=("Helvetica", 15, "bold"),
                 fg=FG, bg=BG_SIDE, pady=8).pack()

        list_frame = tk.Frame(left, bg=BG_SIDE)
        list_frame.pack(fill="both", expand=True, padx=2, pady=2)

        scrollbar = tk.Scrollbar(list_frame)
        scrollbar.pack(side="right", fill="y")

        self.file_listbox = tk.Listbox(
            list_frame, bg=BG_SIDE, fg=FG,
            selectbackground=BG_SEL, selectforeground="#ffffff",
            font=("Helvetica", 13), relief="flat", bd=0,
            activestyle="none", yscrollcommand=scrollbar.set
        )
        self.file_listbox.pack(side="left", fill="both", expand=True)
        scrollbar.config(command=self.file_listbox.yview)
        self.file_listbox.bind("<<ListboxSelect>>", self._on_list_select)

        # Center panel – image viewer
        center = tk.Frame(content, bg=BG)
        center.pack(side="left", fill="both", expand=True)

        # Filename + VA at top
        self.filename_label = tk.Label(center, text="—", font=("Helvetica", 13, "bold"),
                                       fg="#000000", bg=BG)
        self.filename_label.pack(pady=(10, 2))

        self.va_label = tk.Label(center, text="VA: —", font=("Helvetica", 13),
                                 fg="#000000", bg=BG)
        self.va_label.pack(pady=(0, 2))

        self.center_grade_status = tk.Label(center, text="Not graded", font=("Helvetica", 13),
                                            fg="#000000", bg=BG)
        self.center_grade_status.pack(pady=(0, 10))

        self.canvas = tk.Canvas(center, bg="#f0f0f0", highlightthickness=0, cursor="crosshair")
        self.canvas.pack(fill="both", expand=True, padx=10, pady=(0, 4))

        # Zoom controls row
        zoom_row = tk.Frame(center, bg=BG)
        zoom_row.pack()

        zbtn = {"font": ("Helvetica", 13), "relief": "solid", "bd": 1,
                "bg": BG, "fg": FG, "activebackground": BG, "activeforeground": FG,
                "cursor": "hand2", "pady": 3, "padx": 10, "highlightthickness": 0}

        tk.Button(zoom_row, text="−", command=self._zoom_out, **zbtn).pack(side="left", padx=2)
        self.zoom_label = tk.Label(zoom_row, text="100%", font=("Helvetica", 9),
                                   fg=FG_MUTED, bg=BG, width=6)
        self.zoom_label.pack(side="left")
        tk.Button(zoom_row, text="+", command=self._zoom_in, **zbtn).pack(side="left", padx=2)
        tk.Button(zoom_row, text="Reset", command=self._zoom_reset, **zbtn).pack(side="left", padx=(8, 2))

        # Navigation
        nav = tk.Frame(center, bg=BG, pady=6)
        nav.pack()

        nbtn = {"font": ("Helvetica", 13), "relief": "solid", "bd": 1,
                "bg": BG, "fg": FG, "activebackground": BG, "activeforeground": FG,
                "cursor": "hand2", "pady": 4, "padx": 16, "highlightthickness": 0}

        self.prev_btn = tk.Button(nav, text="◀  Prev", command=self._prev_image, **nbtn)
        self.prev_btn.pack(side="left", padx=6)

        self.index_label = tk.Label(nav, text="—", font=("Helvetica", 10),
                                    fg=FG_MUTED, bg=BG, width=12)
        self.index_label.pack(side="left")

        self.next_btn = tk.Button(nav, text="Next  ▶", command=self._next_image, **nbtn)
        self.next_btn.pack(side="left", padx=6)

        # Right panel – grading
        right = tk.Frame(content, bg=BG_SIDE, width=255,
                         highlightthickness=1, highlightbackground=BORDER)
        right.pack(side="right", fill="y")
        right.pack_propagate(False)

        tk.Label(right, text="Image Tools", font=("Helvetica", 13, "bold"),
                 fg=FG, bg=BG_SIDE, pady=10).pack()

        # Brightness control
        brightness_frame = tk.Frame(right, bg=BG_SIDE)
        brightness_frame.pack(fill="x", padx=10, pady=6)

        tk.Label(brightness_frame, text="Brightness:", font=("Helvetica", 11),
                 fg=FG, bg=BG_SIDE).pack(side="left")

        self.brightness_scale = tk.Scale(brightness_frame, from_=0.5, to=2.0, resolution=0.1,
                                         orient="horizontal", bg=BG, fg=FG, length=100,
                                         command=self._on_brightness_change)
        self.brightness_scale.set(1.0)
        self.brightness_scale.pack(side="left", padx=6)

        # Red-free filter button
        self.red_free_btn = tk.Button(right, text="Red-Free Filter", command=self._toggle_red_free,
                                      font=("Helvetica", 11), relief="solid", bd=1,
                                      bg=BG, fg=FG, activebackground=BG, activeforeground=FG,
                                      cursor="hand2", pady=4, padx=12, highlightthickness=0)
        self.red_free_btn.pack(pady=6)

        tk.Frame(right, bg=BORDER, height=1).pack(fill="x", padx=10, pady=10)

        tk.Label(right, text="Grade", font=("Helvetica", 13, "bold"),
                 fg=FG, bg=BG_SIDE, pady=10).pack()

        self.grade_buttons = []
        for grade in DR_GRADES:
            btn_frame = tk.Frame(right, bg=BG_SIDE)
            btn_frame.pack(fill="x", padx=10, pady=2)

            rb = tk.Radiobutton(
                btn_frame,
                text=f"{grade['label']} — {grade['name']}",
                variable=self.selected_grade, value=grade["label"],
                font=("Helvetica", 13, "bold"),
                bg=BG_SIDE, fg=grade["color"],
                selectcolor=BG_SIDE,
                activebackground=BG_SIDE, activeforeground=grade["color"],
                indicatoron=True, relief="flat", bd=0,
                pady=6, padx=4, anchor="w", cursor="hand2",
                command=self._on_grade_selected
            )
            rb.pack(fill="x")
            self.grade_buttons.append(rb)

            tk.Label(btn_frame, text=grade["desc"], font=("Helvetica", 13),
                     fg=FG_MUTED, bg=BG_SIDE, anchor="w",
                     wraplength=210, justify="left").pack(fill="x", padx=20)

        tk.Frame(right, bg=BORDER, height=1).pack(fill="x", padx=10, pady=10)

        tk.Label(right, text="Note", font=("Helvetica", 13, "bold"),
                 fg=FG, bg=BG_SIDE, pady=8).pack()

        self.note_entry = tk.Entry(right, textvariable=self.note_var,
                                   font=("Helvetica", 11), width=28)
        self.note_entry.pack(padx=10, pady=4)
        self.note_entry.bind("<FocusOut>", lambda e: self._on_note_changed())

        hint = ("Keyboard shortcuts:\n"
                "  ← / →    Previous / Next\n"
                "  0 – 4    Grade image\n"
                "  -        Ungradable\n"
                "  scroll   Zoom in / out\n"
                "  drag     Pan image")
        tk.Label(right, text=hint, font=("Helvetica", 13),
                 fg="#a8adb7", bg=BG_SIDE, justify="left",
                 pady=8).pack(side="bottom", padx=12, pady=10)

        # ── Status bar ───────────────────────────────────────────────────
        statusbar = tk.Frame(self, bg=BG_SIDE, pady=4,
                             highlightthickness=1, highlightbackground=BORDER)
        statusbar.pack(fill="x", side="bottom")
        self.status_var = tk.StringVar(value="Open a folder to begin.")
        tk.Label(statusbar, textvariable=self.status_var,
                 font=("Helvetica", 13), fg=FG_MUTED, bg=BG_SIDE).pack(side="left", padx=14)

    # ------------------------------------------------------------------ #
    #  Key & mouse bindings                                                #
    # ------------------------------------------------------------------ #
    def _bind_keys(self):
        self.bind("<Left>",  lambda e: self._prev_image())
        self.bind("<Right>", lambda e: self._next_image())
        for i in range(5):
            self.bind(str(i), lambda e, g=str(i): self._grade_shortcut(g))
        self.bind("-", lambda e: self._grade_shortcut("-1"))

        self.canvas.bind("<MouseWheel>",      self._on_mousewheel)
        self.canvas.bind("<Button-4>",        self._on_scroll_up)
        self.canvas.bind("<Button-5>",        self._on_scroll_down)
        self.canvas.bind("<ButtonPress-1>",   self._on_drag_start)
        self.canvas.bind("<B1-Motion>",       self._on_drag_move)
        self.canvas.bind("<ButtonRelease-1>", self._on_drag_end)
        self.canvas.bind("<Configure>",       self._on_resize)

    def _grade_shortcut(self, grade):
        if self.current_index >= 0:
            self.selected_grade.set(grade)
            self._on_grade_selected()

    # ------------------------------------------------------------------ #
    #  Folder loading                                                      #
    # ------------------------------------------------------------------ #
    def _open_folder(self):
        folder = filedialog.askdirectory(title="Select Image Folder")
        if not folder:
            return
        self.image_folder = folder
        self.image_files = sorted([
            f for f in os.listdir(folder)
            if Path(f).suffix.lower() in SUPPORTED_EXTS
        ])
        if not self.image_files:
            messagebox.showwarning("No Images", "No supported images found in that folder.")
            return

        if not self.use_csv_mode:
            self.grades = {}
            self.file_listbox.delete(0, tk.END)
            for f in self.image_files:
                self.file_listbox.insert(tk.END, f"  {f}")

            short = folder if len(folder) < 50 else "…" + folder[-47:]
            self.folder_label.config(text=short)
            self._set_status(f"Loaded {len(self.image_files)} images from {os.path.basename(folder)}")
            self._load_image(0)
        else:
            self._set_status(f"Folder updated with {len(self.image_files)} images")

    def _open_csv(self):
        xlsx_file = filedialog.askopenfilename(
            title="Select Excel File",
            filetypes=[("Excel files", "*.xlsx"), ("CSV files", "*.csv")],
        )
        if not xlsx_file:
            return

        try:
            if xlsx_file.endswith('.xlsx'):
                self._open_xlsx(xlsx_file)
            else:
                self._open_csv_file(xlsx_file)
        except Exception as ex:
            messagebox.showerror("File Load Error", str(ex))

    def _open_xlsx(self, xlsx_file):
        self.csv_data = {}
        xlsx_images = []

        wb = load_workbook(xlsx_file)
        ws = wb.active

        rows = list(ws.iter_rows(values_only=True))
        if not rows or len(rows[0]) < 2:
            messagebox.showerror("Invalid Excel", "Excel must have at least 2 columns")
            return

        for row in rows[1:]:
            if not row or not row[0]:
                continue
            filename = str(row[0])
            va = str(row[1]) if len(row) > 1 and row[1] else ""
            grader_initial = str(row[2]) if len(row) > 2 and row[2] else ""
            label = str(row[3]) if len(row) > 3 and row[3] else ""
            note = str(row[4]) if len(row) > 4 and row[4] else ""

            self.csv_data[filename] = {
                "va": va,
                "grader_initial": grader_initial,
                "label": label,
                "note": note
            }
            xlsx_images.append(filename)

        # Read grader initial from excel (use first non-empty value)
        for filename in xlsx_images:
            if self.csv_data[filename]["grader_initial"]:
                self.grader_initial_var.set(self.csv_data[filename]["grader_initial"])
                break

        # Check if image folder has been loaded
        if self.image_files:
            # Add missing images to xlsx
            missing_images = [img for img in self.image_files if img not in self.csv_data]
            if missing_images:
                for img in missing_images:
                    self.csv_data[img] = {
                        "va": "",
                        "grader_initial": self.grader_initial_var.get(),
                        "label": "",
                        "note": ""
                    }
                    xlsx_images.append(img)

                # Save updated xlsx
                ws.delete_rows(2, ws.max_row)
                for row_idx, filename in enumerate(xlsx_images, start=2):
                    va = self.csv_data[filename]["va"]
                    grader_initial = self.csv_data[filename]["grader_initial"]
                    label = self.csv_data[filename]["label"]
                    note = self.csv_data[filename]["note"]
                    ws.cell(row=row_idx, column=1).value = filename
                    ws.cell(row=row_idx, column=2).value = va
                    ws.cell(row=row_idx, column=3).value = grader_initial
                    ws.cell(row=row_idx, column=4).value = label
                    ws.cell(row=row_idx, column=5).value = note

                wb.save(xlsx_file)
                messagebox.showinfo("Excel Updated", f"Added {len(missing_images)} missing image(s) to Excel")

            self.image_files = sorted(self.image_files)
        else:
            # No folder loaded yet, use images from excel
            if not xlsx_images:
                messagebox.showwarning("No Images", "No image filenames found in Excel")
                return

            # Ask user to select image folder
            messagebox.showinfo("Select Folder", "Please select the image folder.")
            image_folder = filedialog.askdirectory(title="Select Image Folder")
            if not image_folder:
                return

            self.image_folder = image_folder
            self.image_files = xlsx_images

        self.csv_path = xlsx_file
        self.use_csv_mode = True
        self.grades = {}

        for filename in self.image_files:
            if self.csv_data[filename]["label"]:
                self.grades[filename] = self.csv_data[filename]["label"]

        self.file_listbox.delete(0, tk.END)
        for f in self.image_files:
            self.file_listbox.insert(tk.END, f"  {f}")

        short = xlsx_file if len(xlsx_file) < 50 else "…" + xlsx_file[-47:]
        self.folder_label.config(text=short)
        self._set_status(f"Loaded {len(self.image_files)} images from Excel")
        self._load_image(0)

    def _open_csv_file(self, csv_file):
        self.csv_data = {}
        csv_images = []

        with open(csv_file, "r", newline="", encoding="utf-8") as f:
            reader = csv.reader(f)
            header = next(reader, None)
            if not header or len(header) < 2:
                messagebox.showerror("Invalid CSV", "CSV must have at least 2 columns")
                return

            for row in reader:
                if not row or not row[0]:
                    continue
                filename = row[0]
                va = row[1] if len(row) > 1 else ""
                grader_initial = row[2] if len(row) > 2 else ""
                label = row[3] if len(row) > 3 else ""
                note = row[4] if len(row) > 4 else ""

                self.csv_data[filename] = {
                    "va": va,
                    "grader_initial": grader_initial,
                    "label": label,
                    "note": note
                }
                csv_images.append(filename)

        # Read grader initial from csv (use first non-empty value)
        for filename in csv_images:
            if self.csv_data[filename]["grader_initial"]:
                self.grader_initial_var.set(self.csv_data[filename]["grader_initial"])
                break

        # Check if image folder has been loaded
        if self.image_files:
            # Add missing images to csv
            missing_images = [img for img in self.image_files if img not in self.csv_data]
            if missing_images:
                for img in missing_images:
                    self.csv_data[img] = {
                        "va": "",
                        "grader_initial": self.grader_initial_var.get(),
                        "label": "",
                        "note": ""
                    }
                    csv_images.append(img)

                # Save updated csv
                with open(csv_file, "w", newline="", encoding="utf-8") as f:
                    writer = csv.writer(f)
                    writer.writerow(["filename", "va", "grader_initial", "label", "note"])
                    for filename in csv_images:
                        va = self.csv_data[filename]["va"]
                        grader_initial = self.csv_data[filename]["grader_initial"]
                        label = self.csv_data[filename]["label"]
                        note = self.csv_data[filename]["note"]
                        writer.writerow([filename, va, grader_initial, label, note])

                messagebox.showinfo("CSV Updated", f"Added {len(missing_images)} missing image(s) to CSV")

            self.image_files = sorted(self.image_files)
        else:
            # No folder loaded yet, use images from csv
            if not csv_images:
                messagebox.showwarning("No Images", "No image filenames found in CSV")
                return

            # Ask user to select image folder
            messagebox.showinfo("Select Folder", "Please select the image folder.")
            image_folder = filedialog.askdirectory(title="Select Image Folder")
            if not image_folder:
                return

            self.image_folder = image_folder
            self.image_files = csv_images

        self.csv_path = csv_file
        self.use_csv_mode = True
        self.grades = {}

        for filename in self.image_files:
            if self.csv_data[filename]["label"]:
                self.grades[filename] = self.csv_data[filename]["label"]

        self.file_listbox.delete(0, tk.END)
        for f in self.image_files:
            self.file_listbox.insert(tk.END, f"  {f}")

        short = csv_file if len(csv_file) < 50 else "…" + csv_file[-47:]
        self.folder_label.config(text=short)
        self._set_status(f"Loaded {len(self.image_files)} images from CSV")
        self._load_image(0)

    # ------------------------------------------------------------------ #
    #  Navigation                                                          #
    # ------------------------------------------------------------------ #
    def _load_image(self, index):
        if not self.image_files or index < 0 or index >= len(self.image_files):
            return
        self.current_index = index

        self.file_listbox.selection_clear(0, tk.END)
        self.file_listbox.selection_set(index)
        self.file_listbox.see(index)

        filename = self.image_files[index]
        filepath = os.path.join(self.image_folder, filename)

        try:
            self._pil_image = Image.open(filepath).convert("RGB")
        except Exception as ex:
            self._pil_image = None
            self.canvas.delete("all")
            self.canvas.create_text(400, 300, text=f"Could not load image:\n{ex}",
                                    fill="#cf222e", font=("Helvetica", 12))

        self._zoom  = 1.0
        self._pan_x = 0.0
        self._pan_y = 0.0
        self.brightness = 1.0
        self.red_free_filter = False
        self.brightness_scale.set(1.0)
        self.red_free_btn.config(bg=BG, fg=FG)
        self._render()

        self.filename_label.config(text=filename)

        if self.use_csv_mode and filename in self.csv_data:
            va_text = f"VA: {self.csv_data[filename]['va']}"
        else:
            va_text = "VA: —"
        self.va_label.config(text=va_text)

        self.index_label.config(text=f"{index + 1} / {len(self.image_files)}")

        saved = self.grades.get(filename, "")
        self.selected_grade.set(saved)
        self._refresh_grade_ui(saved)
        self._update_center_grade_status(saved)

        if self.use_csv_mode and filename in self.csv_data:
            self.note_var.set(self.csv_data[filename].get('note', ''))
        else:
            self.note_var.set("")

        self._update_listbox_colors()
        self._update_progress()

    def _prev_image(self):
        if self.current_index > 0:
            self._load_image(self.current_index - 1)

    def _next_image(self):
        if self.current_index < len(self.image_files) - 1:
            self._load_image(self.current_index + 1)

    def _on_list_select(self, event):
        sel = self.file_listbox.curselection()
        if sel:
            self._load_image(sel[0])

    # ------------------------------------------------------------------ #
    #  Zoom / Pan                                                          #
    # ------------------------------------------------------------------ #
    ZOOM_STEP = 1.25
    ZOOM_MIN  = 0.1
    ZOOM_MAX  = 20.0

    def _zoom_in(self):
        self._apply_zoom(self.ZOOM_STEP)

    def _zoom_out(self):
        self._apply_zoom(1.0 / self.ZOOM_STEP)

    def _zoom_reset(self):
        self._zoom  = 1.0
        self._pan_x = 0.0
        self._pan_y = 0.0
        self.brightness = 1.0
        self.red_free_filter = False
        self.brightness_scale.set(1.0)
        self.red_free_btn.config(bg=BG, fg=FG)
        self._render()

    def _apply_zoom(self, factor, cx=None, cy=None):
        if self._pil_image is None:
            return

        new_zoom = max(self.ZOOM_MIN, min(self.ZOOM_MAX, self._zoom * factor))
        if new_zoom == self._zoom:
            return

        cw = self.canvas.winfo_width()  or 700
        ch = self.canvas.winfo_height() or 520
        if cx is None: cx = cw / 2
        if cy is None: cy = ch / 2

        iw, ih = self._pil_image.size
        fit    = min(cw / iw, ch / ih)
        base_w = iw * fit
        base_h = ih * fit

        img_x0 = (cw - base_w) / 2 - self._pan_x * fit * self._zoom
        img_y0 = (ch - base_h) / 2 - self._pan_y * fit * self._zoom

        rel_x = (cx - img_x0) / (fit * self._zoom)
        rel_y = (cy - img_y0) / (fit * self._zoom)

        self._zoom     = new_zoom
        new_img_x0     = cx - rel_x * fit * self._zoom
        new_img_y0     = cy - rel_y * fit * self._zoom
        self._pan_x    = -(new_img_x0 - (cw - base_w) / 2) / (fit * self._zoom)
        self._pan_y    = -(new_img_y0 - (ch - base_h) / 2) / (fit * self._zoom)
        self._render()

    def _on_mousewheel(self, event):
        factor = self.ZOOM_STEP if event.delta > 0 else 1.0 / self.ZOOM_STEP
        self._apply_zoom(factor, cx=event.x, cy=event.y)

    def _on_scroll_up(self, event):
        self._apply_zoom(self.ZOOM_STEP, cx=event.x, cy=event.y)

    def _on_scroll_down(self, event):
        self._apply_zoom(1.0 / self.ZOOM_STEP, cx=event.x, cy=event.y)

    def _on_drag_start(self, event):
        if self._pil_image is None:
            return
        self._drag_start = (event.x, event.y)
        self._drag_pan   = (self._pan_x, self._pan_y)
        self.canvas.config(cursor="fleur")

    def _on_drag_move(self, event):
        if self._drag_start is None or self._pil_image is None:
            return
        dx = event.x - self._drag_start[0]
        dy = event.y - self._drag_start[1]

        iw, ih = self._pil_image.size
        cw = self.canvas.winfo_width()  or 700
        ch = self.canvas.winfo_height() or 520
        fit = min(cw / iw, ch / ih)

        self._pan_x = self._drag_pan[0] - dx / (fit * self._zoom)
        self._pan_y = self._drag_pan[1] - dy / (fit * self._zoom)
        self._render()

    def _on_drag_end(self, event):
        self._drag_start = None
        self._drag_pan   = None
        self.canvas.config(cursor="crosshair")

    # ------------------------------------------------------------------ #
    #  Rendering                                                           #
    # ------------------------------------------------------------------ #
    def _apply_filters(self, image):
        result = image.copy()

        if self.red_free_filter:
            r, g, b = result.split()
            result = Image.merge("RGB", (g, g, b))

        if self.brightness != 1.0:
            enhancer = ImageEnhance.Brightness(result)
            result = enhancer.enhance(self.brightness)

        return result

    def _render(self):
        if self._pil_image is None:
            return

        self.canvas.update_idletasks()
        cw = self.canvas.winfo_width()  or 700
        ch = self.canvas.winfo_height() or 520

        iw, ih = self._pil_image.size
        fit    = min(cw / iw, ch / ih)
        scale  = fit * self._zoom

        vw = cw / scale
        vh = ch / scale

        cx_img = iw / 2 + self._pan_x
        cy_img = ih / 2 + self._pan_y

        left   = int(max(0, cx_img - vw / 2))
        top    = int(max(0, cy_img - vh / 2))
        right  = int(min(iw, cx_img + vw / 2))
        bottom = int(min(ih, cy_img + vh / 2))

        if right <= left or bottom <= top:
            return

        crop  = self._pil_image.crop((left, top, right, bottom))
        crop = self._apply_filters(crop)
        out_w = max(1, int((right - left) * scale))
        out_h = max(1, int((bottom - top) * scale))

        resized = crop.resize((out_w, out_h), Image.LANCZOS)
        self.current_photo = ImageTk.PhotoImage(resized)

        canvas_x = int(left * scale - (cx_img - vw / 2) * scale)
        canvas_y = int(top  * scale - (cy_img - vh / 2) * scale)

        self.canvas.delete("all")
        self.canvas.create_image(canvas_x, canvas_y, anchor="nw", image=self.current_photo)

        self.zoom_label.config(text=f"{int(self._zoom * 100)}%")

    def _on_resize(self, event):
        self._render()

    # ------------------------------------------------------------------ #
    #  Grading                                                             #
    # ------------------------------------------------------------------ #
    def _on_grade_selected(self):
        if self.current_index < 0:
            return
        grade    = self.selected_grade.get()
        filename = self.image_files[self.current_index]
        self.grades[filename] = grade

        if self.use_csv_mode and filename in self.csv_data:
            self.csv_data[filename]["label"] = grade
            self.csv_data[filename]["note"] = self.note_var.get()

        self._refresh_grade_ui(grade)
        self._update_center_grade_status(grade)
        self._update_listbox_colors()
        self._update_progress()

        grade_info = next(g for g in DR_GRADES if g["label"] == grade)
        self._set_status(f"Graded '{filename}' as {grade} – {grade_info['name']}")
        self.after(400, self._auto_advance)

    def _on_grader_changed(self):
        if not self.use_csv_mode:
            return
        grader = self.grader_initial_var.get()
        for filename in self.image_files:
            if filename in self.csv_data:
                self.csv_data[filename]["grader_initial"] = grader

    def _on_note_changed(self):
        if self.current_index < 0 or not self.use_csv_mode:
            return
        filename = self.image_files[self.current_index]
        if filename in self.csv_data:
            self.csv_data[filename]["note"] = self.note_var.get()

    def _auto_advance(self):
        if self.current_index < len(self.image_files) - 1:
            self._next_image()

    def _refresh_grade_ui(self, grade):
        pass

    def _update_center_grade_status(self, grade):
        if grade:
            grade_info = next(g for g in DR_GRADES if g["label"] == grade)
            self.center_grade_status.config(
                text=f"Grade {grade}: {grade_info['name']}",
                fg="#000000"
            )
        else:
            self.center_grade_status.config(text="Not graded", fg="#000000")

    # ------------------------------------------------------------------ #
    #  List coloring / progress                                            #
    # ------------------------------------------------------------------ #
    def _update_listbox_colors(self):
        for i, fname in enumerate(self.image_files):
            if fname in self.grades:
                grade = self.grades[fname]
                color = next(g["color"] for g in DR_GRADES if g["label"] == grade)
                self.file_listbox.itemconfig(i, fg=color)
            else:
                self.file_listbox.itemconfig(i, fg=FG)

    def _update_progress(self):
        self.progress_label.config(
            text=f"{len(self.grades)} / {len(self.image_files)}  graded"
        )

    # ------------------------------------------------------------------ #
    #  CSV export                                                          #
    # ------------------------------------------------------------------ #
    def _save_csv(self):
        if self.use_csv_mode and self.csv_path:
            self._save_csv_to_original()
        else:
            self._save_csv_new()

    def _save_csv_to_original(self):
        if not self.csv_path or not self.csv_data:
            messagebox.showinfo("Nothing to Save", "No file loaded or no data to save.")
            return

        try:
            if self.csv_path.endswith('.xlsx'):
                self._save_xlsx(self.csv_path)
            else:
                self._save_csv_format(self.csv_path)
        except Exception as ex:
            messagebox.showerror("Save Error", str(ex))

    def _save_xlsx(self, xlsx_path):
        wb = load_workbook(xlsx_path)
        ws = wb.active

        ws.cell(row=1, column=1).value = "filename"
        ws.cell(row=1, column=2).value = "va"
        ws.cell(row=1, column=3).value = "grader_initial"
        ws.cell(row=1, column=4).value = "label"
        ws.cell(row=1, column=5).value = "note"

        for row_idx, filename in enumerate(self.image_files, start=2):
            va = self.csv_data.get(filename, {}).get("va", "")
            grader_initial = self.csv_data.get(filename, {}).get("grader_initial", "")
            label = self.csv_data.get(filename, {}).get("label", "")
            note = self.csv_data.get(filename, {}).get("note", "")

            ws.cell(row=row_idx, column=1).value = filename
            ws.cell(row=row_idx, column=2).value = va
            ws.cell(row=row_idx, column=3).value = grader_initial
            ws.cell(row=row_idx, column=4).value = label
            ws.cell(row=row_idx, column=5).value = note

        wb.save(xlsx_path)
        count = len([f for f in self.image_files if self.csv_data.get(f, {}).get("label")])
        self._set_status(f"Saved {count} grade(s) → {xlsx_path}")
        messagebox.showinfo("Saved", f"Saved {count} grade(s) to:\n{xlsx_path}")

    def _save_csv_format(self, csv_path):
        with open(csv_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["filename", "va", "grader_initial", "label", "note"])
            for filename in self.image_files:
                va = self.csv_data.get(filename, {}).get("va", "")
                grader_initial = self.csv_data.get(filename, {}).get("grader_initial", "")
                label = self.csv_data.get(filename, {}).get("label", "")
                note = self.csv_data.get(filename, {}).get("note", "")
                writer.writerow([filename, va, grader_initial, label, note])

        count = len([f for f in self.image_files if self.csv_data.get(f, {}).get("label")])
        self._set_status(f"Saved {count} grade(s) → {csv_path}")
        messagebox.showinfo("Saved", f"Saved {count} grade(s) to:\n{csv_path}")

    def _save_csv_new(self):
        if not self.grades:
            messagebox.showinfo("Nothing to Save", "No images have been graded yet.")
            return

        save_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx"), ("CSV files", "*.csv")],
            initialfile="dr_grades.xlsx",
            title="Save File"
        )
        if not save_path:
            return

        try:
            if save_path.endswith('.xlsx'):
                wb = Workbook()
                ws = wb.active
                ws.cell(row=1, column=1).value = "filename"
                ws.cell(row=1, column=2).value = "label"
                ws.cell(row=1, column=3).value = "note"

                row_idx = 2
                for fname in self.image_files:
                    if fname in self.grades:
                        note = self.csv_data.get(fname, {}).get("note", "") if self.use_csv_mode else ""
                        ws.cell(row=row_idx, column=1).value = fname
                        ws.cell(row=row_idx, column=2).value = self.grades[fname]
                        ws.cell(row=row_idx, column=3).value = note
                        row_idx += 1

                wb.save(save_path)
            else:
                with open(save_path, "w", newline="", encoding="utf-8") as f:
                    writer = csv.writer(f)
                    writer.writerow(["filename", "label", "note"])
                    for fname in self.image_files:
                        if fname in self.grades:
                            note = self.csv_data.get(fname, {}).get("note", "") if self.use_csv_mode else ""
                            writer.writerow([fname, self.grades[fname], note])

            count = len(self.grades)
            self._set_status(f"Saved {count} grade(s) → {save_path}")
            messagebox.showinfo("Saved", f"Saved {count} grade(s) to:\n{save_path}")
        except Exception as ex:
            messagebox.showerror("Save Error", str(ex))

    def _set_status(self, msg):
        self.status_var.set(msg)

    def _on_brightness_change(self, value):
        self.brightness = float(value)
        self._render()

    def _toggle_red_free(self):
        self.red_free_filter = not self.red_free_filter
        if self.red_free_filter:
            self.red_free_btn.config(bg="#e8f5e9", fg="#1b5e20")
        else:
            self.red_free_btn.config(bg=BG, fg=FG)
        self._render()


def main():
    app = DRGrader()
    app.mainloop()


if __name__ == "__main__":
    main()
